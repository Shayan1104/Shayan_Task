import os
import re
import math
import zipfile
import tempfile
import requests
from bs4 import BeautifulSoup
from typing import List, Dict, Any, Optional
from pypdf import PdfReader
from docx import Document as DocxDocument

CHARS_PER_PAGE_ESTIMATE = 2500

class LoadedPage:
    def __init__(self, content: str, source: str, filename: str, page_number: int, total_pages: int, metadata: Optional[Dict[str, Any]] = None):
        self.content = content.strip()
        self.source = source
        self.filename = filename
        self.page_number = page_number
        self.total_pages = total_pages
        self.metadata = metadata or {}

    def to_dict(self) -> Dict[str, Any]:
        return {
            "content": self.content,
            "source": self.source,
            "filename": self.filename,
            "page_number": self.page_number,
            "total_pages": self.total_pages,
            "char_count": len(self.content),
            "metadata": self.metadata
        }


def load_pdf(file_path: str, original_filename: Optional[str] = None) -> List[LoadedPage]:
    """Extracts text page by page from a PDF file."""
    pages: List[LoadedPage] = []
    filename = original_filename or os.path.basename(file_path)
    
    try:
        reader = PdfReader(file_path)
        total = len(reader.pages)
        if total == 0:
            return pages
            
        for idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            pages.append(LoadedPage(
                content=text,
                source=file_path,
                filename=filename,
                page_number=idx + 1,
                total_pages=total,
                metadata={"file_type": "pdf"}
            ))
    except Exception as e:
        print(f"Error loading PDF {file_path}: {e}")
        # Return fallback empty or single page if unreadable
        pages.append(LoadedPage(
            content=f"[Error reading PDF: {str(e)}]",
            source=file_path,
            filename=filename,
            page_number=1,
            total_pages=1,
            metadata={"file_type": "pdf", "error": str(e)}
        ))
    return pages


def load_docx(file_path: str, original_filename: Optional[str] = None) -> List[LoadedPage]:
    """Extracts paragraphs and tables from a DOCX file, estimating pages."""
    filename = original_filename or os.path.basename(file_path)
    try:
        doc = DocxDocument(file_path)
        full_text = []
        for p in doc.paragraphs:
            if p.text.strip():
                full_text.append(p.text)
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    full_text.append(row_text)
                    
        combined = "\n\n".join(full_text)
        return _split_text_into_pages(combined, file_path, filename, "docx")
    except Exception as e:
        print(f"Error loading DOCX {file_path}: {e}")
        return [LoadedPage(f"[Error reading DOCX: {e}]", file_path, filename, 1, 1, {"file_type": "docx"})]


def load_text_file(file_path: str, original_filename: Optional[str] = None) -> List[LoadedPage]:
    """Loads text, markdown, csv, json, code files with encoding fallbacks."""
    filename = original_filename or os.path.basename(file_path)
    encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
    content = ""
    for enc in encodings:
        try:
            with open(file_path, "r", encoding=enc) as f:
                content = f.read()
            break
        except UnicodeDecodeError:
            continue
        except Exception as e:
            content = f"[Error reading file: {e}]"
            break

    ext = os.path.splitext(filename)[1].lower().replace(".", "") or "txt"
    return _split_text_into_pages(content, file_path, filename, ext)


def load_web_url(url: str) -> List[LoadedPage]:
    """WebBaseLoader: fetches and cleans text content from a web URL."""
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    }
    try:
        response = requests.get(url, headers=headers, timeout=15)
        response.raise_for_status()
        
        soup = BeautifulSoup(response.text, "html.parser")
        
        # Remove irrelevant elements
        for element in soup(["script", "style", "nav", "footer", "header", "noscript", "aside", "svg"]):
            element.extract()
            
        title = soup.title.string.strip() if soup.title and soup.title.string else url
        
        # Prefer main content tags if available
        main_content = soup.find("main") or soup.find("article") or soup.find("body") or soup
        text = main_content.get_text(separator="\n", strip=True)
        # Clean extra blank lines
        text = re.sub(r'\n{3,}', '\n\n', text)
        
        doc_name = f"Web: {title[:50]}" if title else f"Web: {url[:50]}"
        pages = _split_text_into_pages(text, url, doc_name, "web")
        for p in pages:
            p.metadata["url"] = url
            p.metadata["title"] = title
        return pages
    except Exception as e:
        print(f"Error fetching web URL {url}: {e}")
        return [LoadedPage(f"[Error fetching URL {url}: {e}]", url, f"Web: {url}", 1, 1, {"file_type": "web", "error": str(e)})]


def _split_text_into_pages(text: str, source: str, filename: str, file_type: str) -> List[LoadedPage]:
    """Splits long text into pages based on estimated characters per page (default ~2500 chars)."""
    text = text.strip()
    if not text:
        return [LoadedPage("", source, filename, 1, 1, {"file_type": file_type})]
        
    total_length = len(text)
    total_pages = max(1, math.ceil(total_length / CHARS_PER_PAGE_ESTIMATE))
    
    pages: List[LoadedPage] = []
    for i in range(total_pages):
        start = i * CHARS_PER_PAGE_ESTIMATE
        end = min((i + 1) * CHARS_PER_PAGE_ESTIMATE, total_length)
        page_text = text[start:end]
        pages.append(LoadedPage(
            content=page_text,
            source=source,
            filename=filename,
            page_number=i + 1,
            total_pages=total_pages,
            metadata={"file_type": file_type}
        ))
    return pages


def load_file_auto(file_path: str, original_filename: Optional[str] = None) -> List[LoadedPage]:
    """Automatically detects file type and routes to the proper loader."""
    name = original_filename or os.path.basename(file_path)
    ext = os.path.splitext(name)[1].lower()
    
    if ext == ".pdf":
        return load_pdf(file_path, name)
    elif ext in [".docx", ".doc"]:
        return load_docx(file_path, name)
    elif ext in [".txt", ".md", ".markdown", ".csv", ".json", ".py", ".js", ".ts", ".html", ".xml", ".yaml", ".yml", ".log", ".rst", ".sql"]:
        return load_text_file(file_path, name)
    elif ext == ".zip":
        return load_zip_file(file_path)
    else:
        # Fallback to text attempt
        return load_text_file(file_path, name)


def load_zip_file(zip_path: str) -> List[LoadedPage]:
    """Extracts a ZIP archive into a temporary folder and loads all contained documents."""
    all_pages: List[LoadedPage] = []
    with tempfile.TemporaryDirectory() as tmp_dir:
        try:
            with zipfile.ZipFile(zip_path, 'r') as zip_ref:
                zip_ref.extractall(tmp_dir)
            all_pages = load_folder(tmp_dir, base_display_name=os.path.basename(zip_path))
        except Exception as e:
            print(f"Error extracting zip {zip_path}: {e}")
    return all_pages


def load_folder(folder_path: str, base_display_name: Optional[str] = None) -> List[LoadedPage]:
    """Recursively traverses a directory and loads all supported documents."""
    all_pages: List[LoadedPage] = []
    if not os.path.exists(folder_path):
        return all_pages

    ignore_dirs = {".git", ".venv", "venv", "__pycache__", "node_modules", ".idea", ".vscode", "chroma_db"}
    supported_exts = {".pdf", ".docx", ".txt", ".md", ".markdown", ".csv", ".json", ".py", ".js", ".ts", ".html", ".xml", ".yaml", ".yml", ".log"}
    
    for root, dirs, files in os.walk(folder_path):
        # Exclude ignored directories in-place
        dirs[:] = [d for d in dirs if d not in ignore_dirs and not d.startswith(".")]
        
        for file in files:
            ext = os.path.splitext(file)[1].lower()
            if ext in supported_exts:
                full_path = os.path.join(root, file)
                rel_path = os.path.relpath(full_path, folder_path)
                display_name = f"{base_display_name}/{rel_path}" if base_display_name else rel_path
                pages = load_file_auto(full_path, original_filename=display_name)
                all_pages.extend(pages)
                
    return all_pages
